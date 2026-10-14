#!/usr/bin/env python3
"""Generate the sample .docx files offered on the gallery cards.

Each template page is parsed and rendered into a Word document, mirroring what
the browser-side exporter (docx-export.js) produces from the same markup, so a
downloaded sample looks like a live download of an unedited template.

Usage:
    python3 tools/generate_sample_docx.py

Requires python-docx (development-only; the website itself ships no Python and
no build step).
"""
from __future__ import annotations

import re
import sys
from html.parser import HTMLParser
from pathlib import Path

try:
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit('python-docx is required: pip install python-docx')

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / 'templates'
OUTPUT = ROOT / 'downloads'

CONTACT_KEYS = ('email', 'phone', 'location', 'website')
VOID_TAGS = {'br', 'img', 'meta', 'link', 'input', 'hr', 'source', 'area',
             'base', 'col', 'embed', 'track', 'wbr'}
SKIP_SUBTREES = {'script', 'style'}


# --------------------------------------------------------------------- DOM --

class Node:
    __slots__ = ('tag', 'attrs', 'children', 'parent', 'text')

    def __init__(self, tag, attrs=None, parent=None, text=None):
        self.tag = tag
        self.attrs = attrs or {}
        self.children = []
        self.parent = parent
        self.text = text

    @property
    def classes(self):
        return (self.attrs.get('class') or '').split()

    def descendants(self):
        for child in self.children:
            yield child
            yield from child.descendants()

    def text_content(self):
        if self.text is not None:
            return self.text
        return ''.join(child.text_content() for child in self.children)

    def closest(self, predicate):
        node = self
        while node is not None:
            if predicate(node):
                return node
            node = node.parent
        return None


class DomParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node('#document')
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = Node(tag, {k: (v or '') for k, v in attrs}, self.stack[-1])
        self.stack[-1].children.append(node)
        if tag not in VOID_TAGS and tag not in SKIP_SUBTREES:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        node = Node(tag, {k: (v or '') for k, v in attrs}, self.stack[-1])
        self.stack[-1].children.append(node)

    def handle_endtag(self, tag):
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                return

    def handle_data(self, data):
        parent = self.stack[-1]
        if parent.tag in SKIP_SUBTREES:
            return
        parent.children.append(Node('#text', parent=parent, text=data))


# ----------------------------------------------------------------- content --

def norm(text):
    return re.sub(r'\s+', ' ', text or '').strip()


def text_of(node):
    return norm(node.text_content())


def lines_of(node):
    """Mirror docx-export.js linesOf(): <br>, real newlines and literal \\n."""
    parts = []

    def walk(current):
        if current.text is not None:
            parts.append(current.text)
            return
        if current.tag == 'br':
            parts.append('\n')
            return
        for child in current.children:
            walk(child)

    walk(node)
    raw = ''.join(parts).replace('\\n', '\n')
    out = []
    for line in raw.split('\n'):
        cleaned = re.sub(r'^[\s\u2022\u00B7*\-]+', '', line).strip()
        if cleaned:
            out.append(cleaned)
    return out


def accent_of(html):
    for name in ('--accent', '--primary'):
        match = re.search(re.escape(name) + r':\s*(#[0-9a-fA-F]{6})', html)
        if match:
            return match.group(1)
    return '#2F5496'


def collect_blocks(html):
    parser = DomParser()
    parser.feed(html)
    root = parser.root

    blocks = []
    contacts = []
    skills = []
    state = {'entry': None, 'current_item': None}

    def flush_contacts():
        if contacts:
            blocks.append(('contact', list(contacts)))
            contacts.clear()

    def flush_skills():
        if skills:
            blocks.append(('para', '  \u00B7  '.join(skills)))
            skills.clear()

    def flush_entry():
        entry = state['entry']
        if not entry:
            return
        label = entry.get('position') or entry.get('degree') or ''
        if label or entry.get('duration'):
            blocks.append(('entry', label, entry.get('duration', '')))
        subtitle = entry.get('company') or entry.get('school')
        if subtitle:
            blocks.append(('subtitle', subtitle))
        for line in entry.get('description', []):
            blocks.append(('bullet', line))
        state['entry'] = None

    def ensure_entry():
        if state['entry'] is None:
            state['entry'] = {'description': []}
        return state['entry']

    nodes = [n for n in root.descendants()
             if 'section-title' in n.classes or 'data-editable' in n.attrs]

    for node in nodes:
        if 'section-title' in node.classes:
            flush_entry()
            flush_skills()
            flush_contacts()
            state['current_item'] = None
            heading = text_of(node)
            if heading:
                blocks.append(('heading', heading))
            continue

        key = node.attrs.get('data-editable')
        if not key or key == 'avatar':
            continue

        if key == 'skill':
            flush_entry()
            flush_contacts()
            state['current_item'] = None
            value = text_of(node)
            if value:
                skills.append(value)
            continue

        item = node.closest(lambda n: 'data-editable-item' in n.attrs)
        if item is not state['current_item']:
            flush_entry()
            state['current_item'] = item

        text = text_of(node)

        if key == 'name':
            flush_skills()
            flush_contacts()
            if text:
                blocks.append(('title', text))
        elif key == 'title':
            flush_skills()
            flush_contacts()
            if text:
                blocks.append(('subtitle-top', text))
        elif key in CONTACT_KEYS:
            flush_entry()
            flush_skills()
            state['current_item'] = None
            if text:
                href = ''
                if key == 'email':
                    href = 'mailto:' + text
                elif key == 'website':
                    href = text if re.match(r'^https?://', text, re.I) else 'https://' + text
                contacts.append((text, href))
        elif key in ('summary', 'language'):
            flush_entry()
            flush_skills()
            flush_contacts()
            state['current_item'] = None
            if text:
                blocks.append(('para', text))
        elif key == 'quote-motto':
            flush_entry()
            flush_skills()
            flush_contacts()
            state['current_item'] = None
            if text:
                blocks.append(('quote', text))
        elif key in ('position', 'degree'):
            flush_skills()
            flush_contacts()
            ensure_entry()[key] = text
        elif key == 'duration':
            if text:
                ensure_entry()['duration'] = text
        elif key in ('company', 'school'):
            if text:
                ensure_entry()[key] = text
        elif key in ('description', 'deal-highlight'):
            ensure_entry()['description'].extend(lines_of(node))
        else:
            if state['current_item'] is not None:
                if text:
                    ensure_entry()['description'].append(text)
            else:
                flush_skills()
                flush_contacts()
                if text:
                    blocks.append(('para', text))

    flush_entry()
    flush_skills()
    flush_contacts()
    return blocks


# -------------------------------------------------------------------- docx --

def add_hyperlink(paragraph, url, text):
    """python-docx has no public hyperlink API, so add the run manually."""
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rel_id)

    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle')
    style.set(qn('w:val'), 'Hyperlink')
    props.append(style)
    run.append(props)

    node = OxmlElement('w:t')
    node.set(qn('xml:space'), 'preserve')
    node.text = text
    run.append(node)
    hyperlink.append(run)

    paragraph._p.append(hyperlink)


def build_document(blocks, title, accent):
    document = docx.Document()

    for section in document.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    heading2 = document.styles['Heading 2']
    heading2.font.name = 'Calibri'
    heading2.font.size = Pt(12)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor.from_string(accent.lstrip('#').upper())
    heading2.paragraph_format.space_before = Pt(14)
    heading2.paragraph_format.space_after = Pt(5)

    bullet = document.styles['List Bullet']
    bullet.font.name = 'Calibri'
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.space_after = Pt(2)

    document.core_properties.title = title
    document.core_properties.author = 'Resume AI'

    for kind, *payload in blocks:
        if kind == 'title':
            paragraph = document.add_heading(payload[0], level=0)
            paragraph.paragraph_format.space_after = Pt(2)
        elif kind == 'subtitle-top':
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(payload[0])
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        elif kind == 'contact':
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            for index, (text, href) in enumerate(payload[0]):
                if index:
                    paragraph.add_run('  \u00B7  ')
                if href:
                    add_hyperlink(paragraph, href, text)
                else:
                    paragraph.add_run(text)
        elif kind == 'heading':
            document.add_heading(payload[0], level=2)
        elif kind == 'entry':
            label, duration = payload[0], payload[1]
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(1)
            if label:
                paragraph.add_run(label).bold = True
            if duration:
                if label:
                    separator = paragraph.add_run('  \u2014  ')
                    separator.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
                run = paragraph.add_run(duration)
                run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        elif kind == 'subtitle':
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(payload[0])
            run.italic = True
            run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        elif kind == 'bullet':
            document.add_paragraph(payload[0], style='List Bullet')
        elif kind == 'quote':
            paragraph = document.add_paragraph(payload[0])
            paragraph.paragraph_format.left_indent = Inches(0.25)
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        elif kind == 'para':
            document.add_paragraph(payload[0])

    return document


def main():
    OUTPUT.mkdir(exist_ok=True)
    count = 0
    for path in sorted(TEMPLATES.glob('*.html')):
        html = path.read_text(encoding='utf-8')
        blocks = collect_blocks(html)
        match = re.search(r'<title>(.*?)</title>', html, re.S)
        title = norm(match.group(1)) if match else path.stem
        document = build_document(blocks, title, accent_of(html))
        target = OUTPUT / f'{path.stem}.docx'
        document.save(str(target))
        count += 1
        print(f'  {path.stem:26} {len(blocks):3} blocks -> {target.name}')
    print(f'\nwrote {count} files to {OUTPUT.relative_to(ROOT)}/')


if __name__ == '__main__':
    main()
